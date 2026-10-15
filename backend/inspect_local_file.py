import os
import openpyxl

path = r"c:\Users\deepak.m_myg\Downloads\All kerala SMS_2025 SEPT TO 2026 AUG STORE DATA.xlsx"
print(f"File exists: {os.path.exists(path)}, size: {os.path.getsize(path)/1024/1024:.2f} MB")

wb = openpyxl.load_workbook(path, read_only=True)
print(f"Sheets: {wb.sheetnames}")
sheet = wb.active
rows = list(sheet.iter_rows(max_row=5, values_only=True))
print(f"First 5 rows: {rows}")
