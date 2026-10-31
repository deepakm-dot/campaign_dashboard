"""
Overall store review: lifetime profile, growth trajectory, latest-quarter vs last year, product mix vs peers,
ranking, trading rhythm, mobile brands, same-town store effect (internal) and festival / GST / market context
(external, each checked against the store's own data where possible). Produces a chat answer + Excel workbook.
"""
import os
import re
import time
import uuid
import calendar
import datetime as dt
from collections import defaultdict

_FIRST_SALE = {'t': 0, 'v': {}}

# External context (public information gathered Oct 2026; NOT stored in the company database)
FESTIVALS = [  # (name, date, days_before_included)
    ('Eid al-Fitr', dt.date(2025, 3, 31), 6), ('Vishu', dt.date(2025, 4, 14), 6),
    ('Onam (Thiruvonam)', dt.date(2025, 9, 5), 9), ('Diwali', dt.date(2025, 10, 20), 9),
    ('Christmas', dt.date(2025, 12, 25), 5), ('Eid al-Fitr', dt.date(2026, 3, 20), 6),
    ('Vishu', dt.date(2026, 4, 14), 6), ('Onam (Thiruvonam)', dt.date(2026, 8, 26), 9),
    ('Diwali', dt.date(2026, 11, 8), 9), ('Christmas', dt.date(2026, 12, 25), 5),
]
GST_DATE = dt.date(2025, 9, 22)  # 56th GST Council: AC, dishwashers, all TVs 28% -> 18%
GST_CATS = ['AIR CONDITIONER', 'TV', 'DISH WASHER']
MARKET_NOTES = [
    "GST 2.0 (effective 22-Sep-2025): GST on air conditioners, dishwashers and all TVs was cut from 28% to 18%, lowering shelf prices and pulling demand forward into the festive season.",
    "India smartphones Q2-2026 (IDC/Omdia): shipments fell about 11-13% YoY while average selling price rose about 14% to a record, driven by the global memory-chip shortage; entry-level segments were hit hardest and buyers moved to mid/premium phones and EMI.",
    "Festival calendar shifts: Thiruvonam was 5-Sep-2025 and 26-Aug-2026 (both inside Jul-Sep); Diwali was 20-Oct-2025 and falls on 8-Nov-2026, so Oct-Dec comparisons will be split differently between October and November.",
]

_RAW_PROD = (
    "if(coalesce(im.product, '') != '', im.product, "
    "multiIf("
    "regexpExtract(s.item_code, '^[A-Z]+', 0) = 'MOB', 'MOBILE', "
    "regexpExtract(s.item_code, '^[A-Z]+', 0) = 'TV', 'TV', "
    "regexpExtract(s.item_code, '^[A-Z]+', 0) = 'LAP', 'LAPTOP', "
    "regexpExtract(s.item_code, '^[A-Z]+', 0) IN ('WSM', 'WM'), 'WASHING MACHINES', "
    "regexpExtract(s.item_code, '^[A-Z]+', 0) = 'REF', 'REFRIGERATORS', "
    "regexpExtract(s.item_code, '^[A-Z]+', 0) = 'TAB', 'TABLET', "
    "regexpExtract(s.item_code, '^[A-Z]+', 0) = 'AC' AND abs(toFloat64(s.sold_price)) >= 1000, 'AIR CONDITIONER', "
    "regexpExtract(s.item_code, '^[A-Z]+', 0) = 'AC' AND abs(toFloat64(s.sold_price)) < 100, 'AC OUTDOOR', "
    "''))"
)
_IS_OD = f"(({_RAW_PROD}) = 'AC OUTDOOR' OR lower(coalesce(im.item_category, '')) = 'ac outdoor' OR lower(coalesce(im.item_name, '')) LIKE 'o/d%')"
PROD = f"if({_IS_OD}, 'AIR CONDITIONER', {_RAW_PROD})"
QTY = f"if({_IS_OD}, 0, if(s.invoice_no LIKE '%-SR-%', -abs(toInt64(s.qty)), if(toFloat64(s.sold_price) < 0, 0, toInt64(s.qty))))"
REV = "toFloat64(s.sold_price) * toInt64(s.qty)"


def _f(x):
    try:
        return float(x)
    except Exception:
        return 0.0


def _d(s):
    return dt.date.fromisoformat(str(s)[:10])


def _pct(new, old):
    return (new - old) / old if old else None


def _fp(new, old):
    p = _pct(new, old)
    return f"{p * 100:+.1f}%" if p is not None else ("new" if new else "-")


def _quarter_label(start):
    names = {1: 'JFM', 4: 'AMJ', 7: 'JAS', 10: 'OND'}
    return f"{names[start.month]} {start.year}"


def _qstart(d):
    return dt.date(d.year, 3 * ((d.month - 1) // 3) + 1, 1)


def _qend(qs):
    m = qs.month + 2
    return dt.date(qs.year, m, calendar.monthrange(qs.year, m)[1])


def _first_sales(ch_query):
    if time.time() - _FIRST_SALE['t'] < 6 * 3600 and _FIRST_SALE['v']:
        return _FIRST_SALE['v']
    # true trading start = first day with >= Rs 1 lakh of sales (ignores 1-invoice test/dummy bills before inauguration)
    rows = ch_query("SELECT branch, toString(min(d)) AS f FROM (SELECT branch, toDate(date) AS d, sum(sold_price) AS r FROM azure_sales_report GROUP BY branch, d HAVING r >= 100000) GROUP BY branch", timeout=120, quote64=True) or []
    _FIRST_SALE['v'] = {r['branch']: _d(r['f']) for r in rows}
    _FIRST_SALE['t'] = time.time()
    return _FIRST_SALE['v']


def store_review(code, name, latest_date_str):
    import ai_service as A
    q = lambda sql, t=120: A.ch_query(sql, timeout=t, quote64=True) or []
    excl = A._EXCL_CATS
    latest = _d(latest_date_str)
    first = _first_sales(A.ch_query).get(code)
    if not first:
        return {"answer": f"No sales found for {name}.", "cards": [], "table": None, "suggestions": []}

    # windows: latest complete calendar quarter and the same quarter a year earlier
    cq = _qstart(latest)
    if _qend(cq) > latest:
        cq = dt.date(cq.year - 1, 10, 1) if cq.month == 1 else dt.date(cq.year, cq.month - 3, 1)
    cq_end = _qend(cq)
    bq = dt.date(cq.year - 1, cq.month, 1)
    bq_end = _qend(bq)
    lab_q, lab_b = _quarter_label(cq), _quarter_label(bq)
    d12 = latest - dt.timedelta(days=364)
    base_ok = first <= bq_end  # store traded in base quarter

    # ---- Q1 store monthly x category ----
    base_filter = "coalesce(im.product, '') != 'SCHEME' AND coalesce(im.item_group, '') != 'SCHEME'"
    rows1 = q(f"""SELECT toString(toStartOfMonth(toDate(s.date))) AS m, {PROD} AS grp, sum({QTY}) AS u, sum({REV}) AS r
                  FROM azure_sales_report s LEFT JOIN item_master im ON im.item_code = s.item_code
                  WHERE s.branch = '{code}' AND toDate(s.date) >= '{first}' AND {base_filter} GROUP BY m, grp""")
    mc = defaultdict(lambda: [0, 0.0])  # (month, cat) -> units, rev
    for r in rows1:
        mc[(_d(r['m']), r['grp'])][0] += int(_f(r['u']))
        mc[(_d(r['m']), r['grp'])][1] += _f(r['r'])
    months = sorted({k[0] for k in mc})
    mon_tot = {m: [sum(v[0] for k, v in mc.items() if k[0] == m), sum(v[1] for k, v in mc.items() if k[0] == m)] for m in months}

    # ---- Q2 store daily ----
    rows2 = q(f"""SELECT toString(toDate(s.date)) AS d, sum({REV}) AS r, sum({QTY}) AS u,
                         uniqExactIf(s.invoice_no, s.sold_price >= 0 AND s.invoice_no NOT LIKE '%-SR-%') AS inv,
                         sumIf(toInt64(s.qty), s.sold_price < 0 OR s.invoice_no LIKE '%-SR-%') AS ret
                  FROM azure_sales_report s LEFT JOIN item_master im ON im.item_code = s.item_code
                  WHERE s.branch = '{code}' AND toDate(s.date) >= '{first}' AND {base_filter} GROUP BY d ORDER BY d""")
    daily = {_d(r['d']): (_f(r['r']), int(_f(r['u'])), int(_f(r['inv'])), int(_f(r['ret']))) for r in rows2}

    # ---- Q3 chain: branch x category windows ----
    w_start = min(bq, d12)
    rows3 = q(f"""SELECT s.branch AS br, any(bm.branch_name) AS nm, {PROD} AS grp,
                         sumIf({REV}, toDate(s.date) >= '{d12}') AS r12, sumIf({QTY}, toDate(s.date) >= '{d12}') AS u12,
                         sumIf({REV}, toDate(s.date) BETWEEN '{cq}' AND '{cq_end}') AS rq,
                         sumIf({REV}, toDate(s.date) BETWEEN '{bq}' AND '{bq_end}') AS rb
                  FROM azure_sales_report s LEFT JOIN item_master im ON im.item_code = s.item_code
                  LEFT JOIN branch_master bm ON bm.code = s.branch
                  WHERE toDate(s.date) >= '{w_start}' AND s.branch NOT IN ('SMC','3GH') AND {base_filter} GROUP BY br, grp""", 180)
    fs = _first_sales(A.ch_query)
    is_future = name.upper().endswith('FUTURE')
    stores = {}
    for r in rows3:
        nm = str(r.get('nm') or r['br'])
        if 'GODOWN' in nm.upper():
            continue
        peer_type = nm.upper().endswith('FUTURE')
        s_ = stores.setdefault(r['br'], {'nm': nm, 'future': peer_type, 'r12': 0.0, 'rq': 0.0, 'rb': 0.0, 'mix': defaultdict(float)})
        s_['r12'] += _f(r['r12']); s_['rq'] += _f(r['rq']); s_['rb'] += _f(r['rb']); s_['mix'][r['grp']] += _f(r['r12'])
    me = stores.get(code)
    peers = {k: v for k, v in stores.items() if v['future'] == is_future and v['r12'] > 0}
    # mix vs peers (share of revenue, last 12 months)
    peer_mix, peer_tot = defaultdict(float), 0.0
    for k, v in peers.items():
        if k == code:
            continue
        peer_tot += v['r12']
        for g, x in v['mix'].items():
            peer_mix[g] += x
    my_tot12 = me['r12'] if me else 0.0

    # ---- Q5 mobile brands ----
    rows5 = q(f"""SELECT coalesce(im.brand,'Other') AS b, sumIf({QTY}, toDate(s.date) BETWEEN '{bq}' AND '{bq_end}') AS ub,
                         sumIf({QTY}, toDate(s.date) BETWEEN '{cq}' AND '{cq_end}') AS uq,
                         sumIf({REV}, toDate(s.date) BETWEEN '{cq}' AND '{cq_end}') AS rq
                  FROM azure_sales_report s LEFT JOIN item_master im ON im.item_code = s.item_code
                  WHERE s.branch = '{code}' AND {PROD} = 'MOBILE' AND toDate(s.date) >= '{bq}' GROUP BY b""")
    brands = sorted([{'b': r['b'], 'ub': int(_f(r['ub'])), 'uq': int(_f(r['uq'])), 'rq': _f(r['rq'])} for r in rows5], key=lambda x: -x['uq'])

    # ---- Q6 same-town store ----
    first_word = name.split()[0].upper()
    sib_rows = q(f"SELECT code, branch_name FROM branch_master WHERE upper(branch_name) LIKE '{first_word}%' AND code != '{code}' AND upper(branch_name) NOT LIKE '%GODOWN%'")
    sib = {}
    if sib_rows:
        codes = ",".join("'" + r['code'] + "'" for r in sib_rows)
        names_s = {r['code']: r['branch_name'] for r in sib_rows}
        for r in q(f"""SELECT s.branch AS br, toString(toStartOfMonth(toDate(s.date))) AS m, sum({REV}) AS r
                       FROM azure_sales_report s LEFT JOIN item_master im ON im.item_code = s.item_code
                       WHERE s.branch IN ({codes}) AND {base_filter} GROUP BY br, m"""):
            sib.setdefault(r['br'], {'nm': names_s.get(r['br'], r['br']), 'm': {}})['m'][_d(r['m'])] = _f(r['r'])

    # ---- Q7 GST window (store + chain) ----
    gst = None
    if first <= GST_DATE - dt.timedelta(days=30) and latest >= GST_DATE + dt.timedelta(days=30):
        g0, g1 = GST_DATE - dt.timedelta(days=30), GST_DATE + dt.timedelta(days=29)
        cats_sql = ",".join("'" + c + "'" for c in GST_CATS)
        def gq(cond):
            return q(f"""SELECT {PROD} AS grp, sumIf({QTY}, toDate(s.date) < '{GST_DATE}') AS pre,
                                sumIf({QTY}, toDate(s.date) >= '{GST_DATE}') AS post
                         FROM azure_sales_report s LEFT JOIN item_master im ON im.item_code = s.item_code
                         WHERE toDate(s.date) BETWEEN '{g0}' AND '{g1}' AND {PROD} IN ({cats_sql}) AND s.branch NOT IN ('SMC','3GH') {cond} GROUP BY grp""")
        gst = {'store': {r['grp']: (int(_f(r['pre'])), int(_f(r['post']))) for r in gq(f"AND s.branch = '{code}'")},
               'chain': {r['grp']: (int(_f(r['pre'])), int(_f(r['post']))) for r in gq("")}}

    # ---- Q8 chain branch x day (benchmarks, launch cohort) and Q9 chain mobile brand mix ----
    rows8 = q(f"""SELECT s.branch AS br, toString(toDate(s.date)) AS d, sum({REV}) AS r, sum({QTY}) AS u,
                         uniqExactIf(s.invoice_no, s.sold_price >= 0 AND s.invoice_no NOT LIKE '%-SR-%') AS inv
                  FROM azure_sales_report s LEFT JOIN item_master im ON im.item_code = s.item_code
                  WHERE toDate(s.date) >= '2025-03-01' AND s.branch NOT IN ('SMC','3GH') AND {base_filter} GROUP BY br, d""", 180)
    type_cond = "upper(bm.branch_name) LIKE '%FUTURE'" if is_future else "upper(bm.branch_name) NOT LIKE '%FUTURE'"
    rows9 = q(f"""SELECT coalesce(im.brand,'Other') AS b, sum({QTY}) AS u
                  FROM azure_sales_report s LEFT JOIN item_master im ON im.item_code = s.item_code
                  LEFT JOIN branch_master bm ON bm.code = s.branch
                  WHERE {PROD} = 'MOBILE' AND toDate(s.date) BETWEEN '{cq}' AND '{cq_end}' AND s.branch NOT IN ('SMC','3GH','{code}')
                        AND {type_cond} AND s.item_code NOT LIKE 'SRV%' GROUP BY b""", 120)

    # =============== analytics ===============
    ins_int, ins_ext, actions, ins_gap = [], [], [], []
    inr = A.format_inr
    life_rev = sum(v[0] for v in daily.values()); life_u = sum(v[1] for v in daily.values())
    life_inv = sum(v[2] for v in daily.values()); life_ret = sum(v[3] for v in daily.values())
    age_m = (latest.year - first.year) * 12 + latest.month - first.month
    last90 = [v for d_, v in daily.items() if d_ > latest - dt.timedelta(days=90)]
    run_rate = sum(v[0] for v in last90) / 90 if last90 else 0
    ins_int.append(f"PROFILE: first sale {first:%d-%b-%Y} ({age_m} months of trading). Lifetime merchandise sales {inr(life_rev)} on {life_u:,} units and {life_inv:,} invoices "
                   f"(average bill {inr(life_rev / life_inv) if life_inv else '-'}, average unit price {inr(life_rev / life_u) if life_u else '-'}). Current run-rate {inr(run_rate)}/day over the last 90 days.")

    # quarters trajectory
    qtab = defaultdict(lambda: [0, 0.0])
    for m, (u, r) in mon_tot.items():
        k = _qstart(m)
        qtab[k][0] += u; qtab[k][1] += r
    qkeys = sorted(qtab)
    full_q = [k for k in qkeys if _qend(k) <= latest]
    if full_q:
        parts = []
        for k in full_q[-6:]:
            parts.append(f"{_quarter_label(k)} {inr(qtab[k][1])} ({qtab[k][0]:,}u)")
        ins_int.append("QUARTERLY TREND: " + "; ".join(parts) + ".")
        if len(full_q) >= 2:
            a, b = full_q[-2], full_q[-1]
            ins_int.append(f"MOMENTUM: {_quarter_label(b)} vs {_quarter_label(a)} revenue {_fp(qtab[b][1], qtab[a][1])}, units {_fp(qtab[b][0], qtab[a][0])}.")
    first_full = [m for m in months if m > first.replace(day=1)]
    if len(first_full) >= 6 and len(months) >= 6:
        early = sum(mon_tot[m][1] for m in first_full[:3]) / 3
        recent = [m for m in months if _qend(_qstart(m)) <= latest or m.month != latest.month][-3:]
        late = sum(mon_tot[m][1] for m in recent) / max(len(recent), 1)
        ins_int.append(f"RAMP-UP: average monthly sales moved from {inr(early)} (first 3 full months) to {inr(late)} (latest 3 complete months), {_fp(late, early)}.")

    # latest quarter vs same quarter last year
    cat_q = defaultdict(lambda: [0, 0, 0.0, 0.0])  # cat -> ub?, uq?, ... [ua, ub, ra, rb]
    for (m, cat), (u, r) in mc.items():
        if bq <= m <= bq_end:
            cat_q[cat][0] += u; cat_q[cat][2] += r
        if cq <= m <= cq_end:
            cat_q[cat][1] += u; cat_q[cat][3] += r
    qa_u = sum(v[0] for v in cat_q.values()); qb_u = sum(v[1] for v in cat_q.values())
    qa_r = sum(v[2] for v in cat_q.values()); qb_r = sum(v[3] for v in cat_q.values())
    if base_ok:
        warn = ""
        months_in_base = (bq.year - first.year) * 12 + bq.month - first.month
        if months_in_base < 6:
            warn = f" Caution: the store opened only {max(months_in_base, 0)} months before {lab_b}, so that base was still ramping up and growth is flattered."
        ins_int.append(f"{lab_q} vs {lab_b}: units {qa_u:,} -> {qb_u:,} ({_fp(qb_u, qa_u)}), revenue {inr(qa_r)} -> {inr(qb_r)} ({_fp(qb_r, qa_r)}).{warn}")
        key_lines = []
        for k in A.KEY_CATS:
            v = cat_q.get(k)
            if v:
                key_lines.append(f"{k} {v[0]:,}->{v[1]:,} units ({_fp(v[1], v[0])}), revenue {_fp(v[3], v[2])}")
        if key_lines:
            ins_int.append("KEY PRODUCTS: " + "; ".join(key_lines) + ".")
        focus = {c: v for c, v in cat_q.items() if c in A.FOCUS_CATS}
        up = sorted([(c, v) for c, v in focus.items() if v[1] > v[0]], key=lambda x: -(x[1][1] - x[1][0]))[:5]
        dn = sorted([(c, v) for c, v in focus.items() if v[1] < v[0]], key=lambda x: x[1][1] - x[1][0])[:5]
        if up:
            ins_int.append("GROWING: " + "; ".join(f"{c} {v[0]:,}->{v[1]:,} ({v[1] - v[0]:+,})" for c, v in up) + ".")
        if dn:
            ins_int.append("DECLINING: " + "; ".join(f"{c} {v[0]:,}->{v[1]:,} ({v[1] - v[0]:+,})" for c, v in dn) + ".")
            for c, v in dn[:3]:
                if c in A.KEY_CATS:
                    actions.append(f"Recover {c}: units fell {v[0]:,} -> {v[1]:,} versus {lab_b}; review stock availability, display and offers for this category.")
    else:
        ins_int.append(f"{lab_b} comparison not possible: the store opened {first:%d-%b-%Y}, after that quarter.")

    # peers / ranking
    rank_txt = ""
    if me:
        eligible = {k: v for k, v in peers.items() if fs.get(k) and fs[k] <= cq}
        if code in eligible or True:
            order = sorted(eligible.items(), key=lambda kv: -kv[1]['rq'])
            pos = next((i for i, (k, _) in enumerate(order) if k == code), None)
            if pos is not None:
                median = sorted(v['rq'] for _, v in order)[len(order) // 2]
                rank_txt = f"{pos + 1} of {len(order)}"
                ins_int.append(f"RANKING: {lab_q} revenue {inr(me['rq'])} ranks {rank_txt} among {'Future' if is_future else 'non-Future'} stores that traded the whole quarter (peer median {inr(median)}; top store {order[0][1]['nm']} {inr(order[0][1]['rq'])}).")
            cohort = {k: v for k, v in eligible.items() if fs.get(k) and abs((fs[k] - first).days) <= 365}
            if len(cohort) > 1 and code in cohort:
                oc = sorted(cohort.items(), key=lambda kv: -kv[1]['rq'])
                pc = next(i for i, (k, _) in enumerate(oc) if k == code)
                ins_int.append(f"AGE-COHORT: among {len(oc)} stores opened within a year of this one, it ranks {pc + 1} (best: {oc[0][1]['nm']} {inr(oc[0][1]['rq'])}).")
    # mix vs peers
    mix_rows = []
    if me and peer_tot > 0 and my_tot12 > 0:
        for g in set(list(peer_mix.keys()) + list(me['mix'].keys())):
            if g not in A.FOCUS_CATS:
                continue
            mine = me['mix'].get(g, 0.0) / my_tot12
            pe = peer_mix.get(g, 0.0) / peer_tot
            mix_rows.append((g, mine, pe, me['mix'].get(g, 0.0), (pe - mine) * my_tot12))
        mix_rows.sort(key=lambda x: -x[3])
        top3 = mix_rows[:3]
        ins_int.append("PRODUCT MIX (last 12 months, share of revenue): " + "; ".join(f"{g} {a * 100:.1f}% (peers {b * 100:.1f}%)" for g, a, b, _, _ in top3) + ".")
        gaps = sorted([x for x in mix_rows if x[4] > 0 and x[2] > 0.01], key=lambda x: -x[4])[:4]
        if gaps:
            ins_int.append("UNDER-INDEXED vs peers (opportunity if the store sold at peer mix): " + "; ".join(f"{g} {a * 100:.1f}% vs {b * 100:.1f}% (~{inr(o)}/yr)" for g, a, b, _, o in gaps) + ".")
            for g, a, b, _, o in gaps[:3]:
                actions.append(f"Lift {g}: share {a * 100:.1f}% vs peer {b * 100:.1f}% - worth about {inr(o)} a year at peer mix.")
        over = sorted([x for x in mix_rows if x[4] < 0 and x[1] > 0.01], key=lambda x: x[4])[:2]
        if over:
            ins_int.append("OVER-INDEXED vs peers (strengths): " + "; ".join(f"{g} {a * 100:.1f}% vs {b * 100:.1f}%" for g, a, b, _, _ in over) + ".")

    # mobile brands
    if brands and sum(b['uq'] for b in brands) > 0:
        tq = sum(b['uq'] for b in brands); tb = sum(b['ub'] for b in brands)
        ins_int.append(f"MOBILE BRANDS {lab_q}: " + "; ".join(f"{b['b']} {b['uq']:,}u ({b['uq'] / tq * 100:.0f}%" + (f", {b['uq'] - b['ub']:+,} vs {lab_b}" if base_ok else "") + ")" for b in brands[:5]) + ".")
        if base_ok and tb:
            losers = sorted([b for b in brands if b['uq'] < b['ub']], key=lambda b: b['uq'] - b['ub'])[:2]
            if losers:
                actions.append("Mobile brand watch: " + ", ".join(f"{b['b']} ({b['ub']:,}->{b['uq']:,})" for b in losers) + " lost units year on year.")

    # rhythm
    recent_d = {d_: v for d_, v in daily.items() if d_ > latest - dt.timedelta(days=180)}
    if len(recent_d) > 30:
        wd = defaultdict(list)
        for d_, v in recent_d.items():
            wd[d_.weekday()].append(v[0])
        avg_all = sum(v[0] for v in recent_d.values()) / len(recent_d)
        names_w = ['Mon', 'Tue', 'Wed', 'Thu', 'Fri', 'Sat', 'Sun']
        idx = {names_w[k]: (sum(v) / len(v)) / avg_all for k, v in wd.items() if v}
        best_w = max(idx, key=idx.get); worst_w = min(idx, key=idx.get)
        ins_int.append(f"WEEKLY RHYTHM (last 180 days): strongest {best_w} ({idx[best_w] * 100:.0f}% of average day), weakest {worst_w} ({idx[worst_w] * 100:.0f}%).")
        actions.append(f"Plan promotions/staffing around the weak {worst_w} and the strong {best_w}.")
    l30 = [v[0] for d_, v in daily.items() if d_ > latest - dt.timedelta(days=30)]
    p30 = [v[0] for d_, v in daily.items() if latest - dt.timedelta(days=60) < d_ <= latest - dt.timedelta(days=30)]
    if l30 and p30:
        ins_int.append(f"LAST 30 DAYS: {inr(sum(l30))} vs previous 30 days {inr(sum(p30))} ({_fp(sum(l30), sum(p30))}).")
    zero_days = [d_ for d_ in (first + dt.timedelta(days=i) for i in range((latest - first).days)) if d_ not in daily and d_ > latest - dt.timedelta(days=180)]
    if zero_days:
        ins_int.append(f"NO-SALE DAYS in last 180 days: {len(zero_days)} (e.g. {', '.join(d_.strftime('%d-%b') for d_ in zero_days[:5])}) - check for closures or data gaps.")
    if life_u:
        ins_int.append(f"RETURNS: {life_ret:,} units returned against {life_u + life_ret:,} sold lifetime ({life_ret / (life_u + life_ret) * 100:.1f}%).")
    # same-town store
    for sc, sv in sib.items():
        ms = sv['m']
        if not ms:
            continue
        s_first = min(ms)
        if s_first < first.replace(day=1):
            pre = [ms.get(first.replace(day=1) - dt.timedelta(days=30 * i), None) for i in range(1, 7)]
            pre = [x for x in (ms.get(dt.date((first.replace(day=1) - dt.timedelta(days=30 * i)).year, (first.replace(day=1) - dt.timedelta(days=30 * i)).month, 1)) for i in range(1, 7)) if x]
            comp = sorted(m for m in ms if m <= latest.replace(day=1) - dt.timedelta(days=1))[-6:]
            if pre and comp:
                pa = sum(pre) / len(pre); pb = sum(ms[m] for m in comp) / len(comp)
                me_recent = sum(mon_tot[m][1] for m in months if m in comp) / max(len([m for m in months if m in comp]), 1)
                ins_int.append(f"SAME-TOWN STORE {sv['nm']}: average monthly sales {inr(pa)} in the 6 months before {name} opened vs {inr(pb)} in the latest 6 complete months ({_fp(pb, pa)}); "
                               f"together with {name} ({inr(me_recent)}/month) the town earns {inr(pb + me_recent)}/month. This is correlation, not proof of cannibalisation.")

    # ---------- potential vs performance (benchmarks against peer stores) ----------
    D = defaultdict(dict)
    for r in rows8:
        b_ = r['br']; dd_ = _d(r['d'])
        if b_ not in stores or not fs.get(b_) or dd_ < fs[b_]:
            continue
        D[b_][dd_] = (_f(r['r']), int(_f(r['u'])), int(_f(r['inv'])))
    fut = [b_ for b_ in D if stores[b_]['future'] == is_future]
    kind = 'Future' if is_future else 'non-Future'

    def _rv(b_, d0, d1, i=0):
        return sum(v[i] for dd_, v in D[b_].items() if d0 <= dd_ <= d1)

    def _qt(vals, p):
        s_ = sorted(vals)
        return s_[int(round(p * (len(s_) - 1)))] if s_ else 0
    td = dt.timedelta
    rank_rows, launch_rows, bench_rows, catgap_rows, brand_gap = [], [], [], [], []
    last_rank = open_rank = None
    if code in D:
        # 1. quarter-by-quarter rank among stores that traded the whole quarter
        qs_ = _qstart(first)
        while _qend(qs_) <= latest:
            qe_ = _qend(qs_); mine_ = _rv(code, qs_, qe_)
            pv = [_rv(b_, qs_, qe_) for b_ in fut if b_ != code and fs[b_] <= qs_]
            if pv and mine_ > 0:
                rank_rows.append((qs_, 1 + sum(1 for v in pv if v > mine_), len(pv) + 1, mine_, _qt(pv, .5), _qt(pv, .75), max(pv)))
            qs_ = dt.date(qs_.year + (qs_.month + 3 > 12), (qs_.month + 2) % 12 + 1, 1)
        if rank_rows:
            ins_gap.append("RANK TREND (quarterly revenue rank among " + kind + " stores that traded the whole quarter): " +
                           "; ".join(f"{_quarter_label(a)} #{rk}/{n} ({m_ / md:.2f}x median)" for a, rk, n, m_, md, _, _ in rank_rows if md) + ".")
            last_rank = rank_rows[-1]
        # 2. launch strength vs other new stores, and how well the launch level was held
        d1_ = daily.get(first, (0, 0, 0, 0))[0]
        wk1 = _rv(code, first, first + td(days=6)); m1 = _rv(code, first, first + td(days=29))
        for b_ in fut:
            if fs[b_] < dt.date(2025, 3, 1) or fs[b_] + td(days=30) > latest:
                continue
            o_ = _rv(b_, fs[b_], fs[b_] + td(days=29))
            if o_ <= 0:
                continue
            m6_ = _rv(b_, fs[b_] + td(days=150), fs[b_] + td(days=179)) if fs[b_] + td(days=180) <= latest else None
            cur_ = _rv(b_, latest - td(days=29), latest) if fs[b_] + td(days=60) <= latest else None
            launch_rows.append((b_, stores[b_]['nm'], fs[b_], o_, m6_, cur_))
        me_l = next((x for x in launch_rows if x[0] == code), None)
        if me_l:
            lo = sorted(launch_rows, key=lambda x: -x[3])
            open_rank = next(i for i, x in enumerate(lo) if x[0] == code) + 1
            r6 = [x[4] / x[3] for x in launch_rows if x[4] is not None]
            rc = [x[5] / x[3] for x in launch_rows if x[5] is not None]
            txt = (f"LAUNCH (inaugurated {first:%d-%b-%Y}): day-1 sales {inr(d1_)}, first week {inr(wk1)}, first 30 days {inr(m1)} - rank {open_rank} of {len(lo)} "
                   f"{kind} stores opened since Mar-2025 (best launch: {lo[0][1]} {inr(lo[0][3])}).")
            ins_gap.append(txt)
            if me_l[4] is not None and r6:
                ins_gap.append(f"LAUNCH HOLD: sales in month 6 were {me_l[4] / me_l[3] * 100:.0f}% of the opening 30 days (new-store median {_qt(r6, .5) * 100:.0f}%)" +
                               (f"; the last 30 days are {me_l[5] / me_l[3] * 100:.0f}% of the opening 30 days (median {_qt(rc, .5) * 100:.0f}%)." if me_l[5] is not None and rc else "."))
        # 3. traffic vs ticket benchmark, last 90 days
        d90 = latest - td(days=89)
        for b_ in fut:
            if fs[b_] > d90:
                continue
            r_, u_, i_ = _rv(b_, d90, latest, 0), _rv(b_, d90, latest, 1), _rv(b_, d90, latest, 2)
            if i_ > 0 and u_ > 0:
                bench_rows.append((b_, stores[b_]['nm'], r_ / 90, i_ / 90, r_ / i_, u_ / i_, r_ / u_))
        me_b = next((x for x in bench_rows if x[0] == code), None)
        oth = [x for x in bench_rows if x[0] != code]
        if me_b and len(oth) >= 5:
            def rk_(i):
                return 1 + sum(1 for x in oth if x[i] > me_b[i])
            parts = []
            for i, lab, fmt in ((2, 'revenue/day', inr), (3, 'bills/day', lambda v: f"{v:,.0f}"), (4, 'average bill', inr), (5, 'units per bill', lambda v: f"{v:.2f}"), (6, 'average unit price', inr)):
                vals = [x[i] for x in oth]
                parts.append(f"{lab} {fmt(me_b[i])} (#{rk_(i)}/{len(oth) + 1}; median {fmt(_qt(vals, .5))}, top-quartile {fmt(_qt(vals, .75))})")
            ins_gap.append(f"BENCHMARK (last 90 days vs {len(oth)} {kind} stores that traded all 90 days): " + "; ".join(parts) + ".")
            p75 = {i: _qt([x[i] for x in oth], .75) for i in (2, 3, 4)}; p50r = _qt([x[2] for x in oth], .5)
            gap75 = max(0, p75[2] - me_b[2]) * 365; gap50 = max(0, p50r - me_b[2]) * 365
            traffic = max(0, p75[3] - me_b[3]) * me_b[4] * 365; ticket = max(0, p75[4] - me_b[4]) * me_b[3] * 365
            ins_gap.append(f"GAP TO PEERS: at the current run-rate the store earns {inr(me_b[2] * 365)} a year; matching the peer median adds about {inr(gap50)} and matching the top quartile about {inr(gap75)}.")
            if traffic or ticket:
                lever = 'footfall / conversion (number of bills)' if traffic >= ticket else 'ticket size (value per bill)'
                ins_gap.append(f"WHICH LEVER: closing the bill-count gap to top-quartile is worth ~{inr(traffic)}/yr and the bill-value gap ~{inr(ticket)}/yr (approximate, treated separately) - the bigger lever is {lever}.")
                actions.append(f"Main lever is {lever}: " + ("drive walk-ins and conversion (local marketing, lead follow-up, sales-staff incentives, in-store offers)." if traffic >= ticket else "raise basket value (bundles, EMI/exchange, premium-model push, accessory attach)."))
        # 4. category gap vs top-quartile peers (last 12 months, absolute revenue)
        elig12 = [k for k, v in peers.items() if k != code and fs.get(k) and fs[k] <= d12 and k in D]
        if me and len(elig12) >= 5:
            for g in A.FOCUS_CATS:
                vals = [peers[k]['mix'].get(g, 0.0) for k in elig12]
                p75c, p50c = _qt(vals, .75), _qt(vals, .5)
                mine_c = me['mix'].get(g, 0.0)
                catgap_rows.append((g, mine_c, p50c, p75c, max(0, p75c - mine_c), 1 + sum(1 for v in vals if v > mine_c)))
            catgap_rows.sort(key=lambda x: -x[4])
            tops = [x for x in catgap_rows if x[4] > 0][:5]
            if tops:
                ins_gap.append(f"CATEGORY GAP (last 12 months revenue vs top-quartile of {len(elig12)} {kind} stores): " +
                               "; ".join(f"{g} {inr(m_)} vs {inr(p75_)} (gap {inr(gp)}, rank {rk}/{len(elig12) + 1})" for g, m_, _, p75_, gp, rk in tops) + ".")
                for g, m_, _, p75_, gp, rk in tops[:3]:
                    actions.append(f"Close the {g} gap: {inr(m_)} vs top-quartile {inr(p75_)} over 12 months (rank {rk}/{len(elig12) + 1}) - about {inr(gp)} a year.")
        # 5. mobile brand mix vs chain
        ch_tot = sum(int(_f(r['u'])) for r in rows9)
        my_tot_b = sum(b['uq'] for b in brands)
        if ch_tot > 0 and my_tot_b > 0:
            chb = {r['b']: int(_f(r['u'])) / ch_tot for r in rows9}
            mine_b = {b['b']: b['uq'] / my_tot_b for b in brands}
            for br_, sh in chb.items():
                if sh >= 0.03:
                    brand_gap.append((br_, mine_b.get(br_, 0.0), sh, (sh - mine_b.get(br_, 0.0)) * my_tot_b))
            brand_gap.sort(key=lambda x: -x[3])
            bg = [x for x in brand_gap if x[3] > 0][:3]
            if bg:
                ins_gap.append(f"MOBILE BRAND GAP ({lab_q}, share of units vs other {kind} stores): " + "; ".join(f"{b} {a * 100:.0f}% vs {s * 100:.0f}% (~{g_:,.0f} units a quarter)" for b, a, s, g_ in bg) + ".")
                for b, a, s, g_ in bg[:2]:
                    actions.append(f"Mobile brand {b}: share {a * 100:.0f}% vs {s * 100:.0f}% at other {kind} stores - check stock depth, display and brand-promoter support (~{g_:,.0f} units a quarter).")
        # verdict
        if last_rank and open_rank:
            a_, rk_, n_, m_, md_, _, mx_ = last_rank
            ins_gap.insert(0, f"VERDICT: the store launched at rank {open_rank} of {len(launch_rows)} new {kind} stores but ranks {rk_} of {n_} {kind} stores in {_quarter_label(a_)} "
                              f"({m_ / md_:.2f}x the peer median; the best store sells {mx_ / m_:.1f}x as much). " +
                              ("The launch strength has not been converted into lasting performance - see the gaps below." if rk_ > open_rank else "Performance has kept pace with the launch."))

    # ---------- external ----------
    fest_rows = []
    for fname, fdate, pre in FESTIVALS:
        w0 = fdate - dt.timedelta(days=pre)
        if fdate > latest or first > w0 - dt.timedelta(days=30):
            continue
        win = [daily[d_][0] for d_ in (w0 + dt.timedelta(days=i) for i in range(pre + 1)) if d_ in daily]
        base = [v[0] for d_, v in daily.items() if fdate - dt.timedelta(days=60) <= d_ <= fdate + dt.timedelta(days=30) and not (w0 <= d_ <= fdate)]
        if len(win) >= max(pre - 1, 3) and len(base) >= 20:
            lift = (sum(win) / len(win)) / (sum(base) / len(base))
            fest_rows.append((fname, fdate, lift, sum(win) / len(win), sum(base) / len(base)))
    if fest_rows:
        ins_ext.append("FESTIVAL LIFT (average daily sales in the run-up window vs the surrounding 90 days, this store's own data): " +
                       "; ".join(f"{n} {d_:%Y} {l:.2f}x" for n, d_, l, _, _ in fest_rows) + ".")
        best = max(fest_rows, key=lambda x: x[2])
        ins_ext.append(f"The store responds most to {best[0]} {best[1]:%Y} ({best[2]:.2f}x). Upcoming Diwali is on 08-Nov-2026 - stock and staffing should peak in the 9 days before it.")
        actions.append(f"Prepare for Diwali (08-Nov-2026): this store's best festival lift so far is {best[0]} {best[1]:%Y} at {best[2]:.2f}x average days.")
    if gst:
        parts = []
        for c in GST_CATS:
            s_pre, s_post = gst['store'].get(c, (0, 0)); c_pre, c_post = gst['chain'].get(c, (0, 0))
            if s_pre + s_post:
                parts.append(f"{c} {s_pre}->{s_post} units ({_fp(s_post, s_pre)}; chain {_fp(c_post, c_pre)})")
        if parts:
            ins_ext.append("GST 2.0 CHECK (30 days before vs after 22-Sep-2025, store vs whole chain): " + "; ".join(parts) + ". Onam (05-Sep-2025) and the festive build-up overlap this window, so treat it as indicative.")
    mob_u = cat_q.get('MOBILE')
    if mob_u and base_ok and mob_u[0]:
        asp_a = mob_u[2] / mob_u[0]; asp_b = mob_u[3] / mob_u[1] if mob_u[1] else 0
        ins_ext.append(f"MARKET CONTEXT - SMARTPHONES: store mobile units {_fp(mob_u[1], mob_u[0])} and average selling price {inr(asp_a)} -> {inr(asp_b)} ({_fp(asp_b, asp_a)}) "
                       f"{lab_b} to {lab_q}, against a national slowdown in shipments with record prices (see notes).")
    ins_ext.extend(["NOTE: " + n for n in MARKET_NOTES])
    if name.upper().startswith('THIRUVALLA'):
        ins_ext.append("CATCHMENT POTENTIAL: Thiruvalla / Pathanamthitta is widely cited as a remittance-led, high-deposit market (large NRI and healthcare-migrant base), so purchasing power is high and skews to premium products and EMI - this supports the 'high potential' view; the same sources note an ageing resident population and sensitivity to Gulf/overseas job and exchange-rate shifts. Public-source context only; no footfall or competitor data is held in the database.")
    ins_ext.append("External items are public-source context gathered in Oct 2026; the dashboard database has no competitor, footfall, weather or local-economy data, so these cannot be verified from internal records.")

    # ---------- Excel ----------
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb = Workbook(); hdr = PatternFill('solid', fgColor='1F2A56'); keyf = PatternFill('solid', fgColor='FFF2CC')
    red, green = Font(color='C00000', bold=True), Font(color='2E7D32', bold=True)

    def sheet(title, headers, data, num=(), pct=(), delta=(), keycol=None):
        ws = wb.create_sheet(title)
        ws.append(headers)
        for c in ws[1]:
            c.font = Font(bold=True, color='FFFFFF'); c.fill = hdr; c.alignment = Alignment(horizontal='center', wrap_text=True)
        for r in data:
            ws.append(r)
        for i in range(2, ws.max_row + 1):
            for ci in num: ws.cell(i, ci).number_format = '#,##0'
            for ci in pct: ws.cell(i, ci).number_format = '0.0%'
            for ci in delta:
                v = ws.cell(i, ci).value
                if isinstance(v, (int, float)) and v != 0:
                    ws.cell(i, ci).font = red if v < 0 else green
            if keycol and ws.cell(i, keycol).value in A.KEY_CATS:
                for cc in range(1, len(headers) + 1): ws.cell(i, cc).fill = keyf
        for ci, h in enumerate(headers, 1):
            w = max([len(str(h))] + [len(str(ws.cell(i, ci).value or '')) for i in range(2, min(ws.max_row, 120) + 1)])
            ws.column_dimensions[get_column_letter(ci)].width = min(max(w + 2, 11), 46)
        ws.freeze_panes = 'A2'
        return ws

    w0 = wb.active; w0.title = 'Review'
    w0.append([f"Overall store review: {name}"]); w0['A1'].font = Font(bold=True, size=14)
    w0.append([f"Data to {latest:%d-%b-%Y}. Total store sales (inter-branch schemes & demo excluded); units net of returns."])
    w0.append([])
    if ins_gap:
        w0.append(["POTENTIAL VS PERFORMANCE DIAGNOSIS"]); w0.cell(w0.max_row, 1).font = Font(bold=True, color='C00000')
        for l in ins_gap: w0.append([l])
        w0.append([])
    w0.append(["INTERNAL INSIGHTS"]); w0.cell(w0.max_row, 1).font = Font(bold=True, color='1F2A56')
    for l in ins_int: w0.append([l])
    w0.append([]); w0.append(["EXTERNAL / MARKET INSIGHTS"]); w0.cell(w0.max_row, 1).font = Font(bold=True, color='1F2A56')
    for l in ins_ext: w0.append([l])
    w0.append([]); w0.append(["SUGGESTED ACTIONS"]); w0.cell(w0.max_row, 1).font = Font(bold=True, color='1F2A56')
    for l in (actions or ["No critical gaps detected from the data."]): w0.append([l])
    w0.column_dimensions['A'].width = 170
    for r in w0.iter_rows(min_row=4):
        r[0].alignment = Alignment(wrap_text=True, vertical='top')

    sheet('Monthly', ['Month', 'Units', 'Revenue', 'Rev MoM %'],
          [[f"{m:%b %Y}", mon_tot[m][0], mon_tot[m][1], _pct(mon_tot[m][1], mon_tot[months[i - 1]][1]) if i else None] for i, m in enumerate(months)], num=(2, 3), pct=(4,))
    sheet('Quarterly', ['Quarter', 'Units', 'Revenue', 'Rev QoQ %', 'Complete?'],
          [[_quarter_label(k), qtab[k][0], qtab[k][1], _pct(qtab[k][1], qtab[qkeys[i - 1]][1]) if i else None, 'Yes' if _qend(k) <= latest else 'No (in progress)'] for i, k in enumerate(qkeys)], num=(2, 3), pct=(4,))
    fy = defaultdict(lambda: [0, 0.0])
    for m, (u, r) in mon_tot.items():
        k = m.year if m.month >= 4 else m.year - 1
        fy[k][0] += u; fy[k][1] += r
    sheet('Financial Year', ['FY (Apr-Mar)', 'Units', 'Revenue'], [[f"FY {k}-{str(k + 1)[2:]}", fy[k][0], fy[k][1]] for k in sorted(fy)], num=(2, 3))
    allcats = sorted({k[1] for k in mc}, key=lambda c: (0 if c in A.KEY_CATS else 1, -sum(v[1] for kk, v in mc.items() if kk[1] == c)))
    sheet(f'{lab_q} vs {lab_b}', ['Category', f'Units {lab_b}', f'Units {lab_q}', 'Unit Change', 'Unit %', f'Revenue {lab_b}', f'Revenue {lab_q}', 'Revenue Change'],
          [[c, cat_q[c][0], cat_q[c][1], cat_q[c][1] - cat_q[c][0], _pct(cat_q[c][1], cat_q[c][0]), cat_q[c][2], cat_q[c][3], cat_q[c][3] - cat_q[c][2]] for c in allcats if c in cat_q],
          num=(2, 3, 4, 6, 7, 8), pct=(5,), delta=(4, 8), keycol=1)
    sheet('Mix vs Peers', ['Category', 'Store share', 'Peer share', 'Gap (pts)', 'Store revenue 12m', 'Revenue opportunity at peer mix'],
          [[g, a, b, a - b, r_, o] for g, a, b, r_, o in sorted(mix_rows, key=lambda x: -x[4])], num=(5, 6), pct=(2, 3, 4), delta=(6,), keycol=1)
    sheet('Mobile Brands', ['Brand', f'Units {lab_b}', f'Units {lab_q}', 'Unit Change', f'Revenue {lab_q}'],
          [[b['b'], b['ub'], b['uq'], b['uq'] - b['ub'], b['rq']] for b in brands], num=(2, 3, 4, 5), delta=(4,))
    sheet('Festival Lift', ['Festival', 'Date', 'Avg daily sales in window', 'Avg daily sales baseline', 'Lift (x)'],
          [[n, d_.strftime('%d-%b-%Y'), a, b, round(l, 2)] for n, d_, l, a, b in fest_rows], num=(3, 4))
    if gst:
        sheet('GST Impact', ['Category', 'Store 30d before', 'Store 30d after', 'Store change', 'Chain 30d before', 'Chain 30d after', 'Chain change'],
              [[c, gst['store'].get(c, (0, 0))[0], gst['store'].get(c, (0, 0))[1], _pct(gst['store'].get(c, (0, 0))[1], gst['store'].get(c, (0, 0))[0]),
                gst['chain'].get(c, (0, 0))[0], gst['chain'].get(c, (0, 0))[1], _pct(gst['chain'].get(c, (0, 0))[1], gst['chain'].get(c, (0, 0))[0])] for c in GST_CATS],
              num=(2, 3, 5, 6), pct=(4, 7))
    if sib:
        allm = sorted({m for v in sib.values() for m in v['m']} | set(months))
        sheet('Same-Town Stores', ['Month', name] + [v['nm'] for v in sib.values()],
              [[f"{m:%b %Y}", mon_tot.get(m, [0, 0])[1]] + [v['m'].get(m, 0) for v in sib.values()] for m in allm], num=tuple(range(2, 3 + len(sib))))
    if rank_rows:
        sheet('Rank Trend', ['Quarter', 'Rank', 'Stores compared', 'Store revenue', 'Peer median', 'Peer top-quartile', 'Best store', 'x Median'],
              [[_quarter_label(a), rk, n, m_, md, p75_, mx_, round(m_ / md, 2) if md else None] for a, rk, n, m_, md, p75_, mx_ in rank_rows], num=(4, 5, 6, 7))
    if launch_rows:
        sheet('Launch Cohort', ['Store', 'Opened', 'First 30 days', 'Month-6 (30d)', 'Last 30 days', 'Month-6 % of opening', 'Last 30d % of opening'],
              [[nm_, o_.strftime('%d-%b-%Y'), v, m6, cur, (m6 / v) if m6 is not None and v else None, (cur / v) if cur is not None and v else None]
               for _, nm_, o_, v, m6, cur in sorted(launch_rows, key=lambda x: -x[3])], num=(3, 4, 5), pct=(6, 7))
    if bench_rows:
        sheet('Peer Benchmark', ['Store', 'Revenue/day (90d)', 'Bills/day', 'Average bill', 'Units per bill', 'Avg unit price'],
              [[x[1], x[2], x[3], x[4], round(x[5], 2), x[6]] for x in sorted(bench_rows, key=lambda x: -x[2])], num=(2, 3, 4, 6))
    if catgap_rows:
        sheet('Category Gap', ['Category', 'Store revenue 12m', 'Peer median', 'Peer top-quartile', 'Gap to top-quartile', 'Rank'],
              [[g, m_, p50_, p75_, gp, rk] for g, m_, p50_, p75_, gp, rk in catgap_rows], num=(2, 3, 4, 5), keycol=1)
    if brand_gap:
        sheet('Mobile Brand Gap', ['Brand', 'Store share', 'Peer-chain share', 'Units/quarter gap'],
              [[b, a, s, g_] for b, a, s, g_ in brand_gap], num=(4,), pct=(2, 3), delta=(4,))
    order_all = sorted(peers.items(), key=lambda kv: -kv[1]['rq'])
    sheet('Peer Ranking', ['Rank', 'Store', f'Revenue {lab_q}', f'Revenue {lab_b}', 'YoY %', 'Opened'],
          [[i + 1, v['nm'], v['rq'], v['rb'], _pct(v['rq'], v['rb']) if v['rb'] else None, fs.get(k).strftime('%d-%b-%Y') if fs.get(k) else ''] for i, (k, v) in enumerate(order_all)], num=(3, 4), pct=(5,))
    wsd = sheet('Daily', ['Date', 'Revenue', 'Units', 'Invoices', 'Returned units'], [[d_.strftime('%Y-%m-%d'), v[0], v[1], v[2], v[3]] for d_, v in sorted(daily.items())], num=(2, 3, 4, 5))

    os.makedirs(A.EXPORT_DIR, exist_ok=True)
    fname = f"insights_store_review_{re.sub(r'[^A-Za-z0-9]+', '_', name).strip('_')}_{uuid.uuid4().hex[:10]}.xlsx"
    wb.save(os.path.join(A.EXPORT_DIR, fname))

    # ---------- chat response ----------
    answer = (f"**Overall review: {name}** (data to {latest:%d-%b-%Y})\n\n" +
              (("**Potential vs performance diagnosis**\n" + "\n".join(f"- {l}" for l in ins_gap) + "\n\n") if ins_gap else "") +
              "**Internal insights**\n" + "\n".join(f"- {l}" for l in ins_int) +
              "\n\n**External / market insights**\n" + "\n".join(f"- {l}" for l in ins_ext) +
              "\n\n**Suggested actions**\n" + "\n".join(f"- {l}" for l in (actions or ["No critical gaps detected from the data."])))
    qtbl_rows = [[_quarter_label(k), f"{qtab[k][0]:,}", A.format_inr_exact(qtab[k][1]), (_fp(qtab[k][1], qtab[qkeys[i - 1]][1]) if (i and qtab[qkeys[i - 1]][1] >= 1e6 and _qend(k) <= latest) else '-'), 'Yes' if _qend(k) <= latest else 'In progress'] for i, k in enumerate(qkeys)]
    cards = [{"label": "First Sale", "value": f"{first:%d %b %Y}", "color": "#6b7280"},
             {"label": "Lifetime Sales", "value": inr(life_rev), "color": "#059669"},
             {"label": f"{lab_q} Revenue", "value": inr(qb_r), "color": "#2563eb"},
             {"label": f"Rank ({'Future' if is_future else 'peers'})", "value": rank_txt or '-', "color": "#d97706"}]
    return {"answer": answer, "cards": cards,
            "table": {"headers": ["Quarter", "Units", "Revenue", "QoQ", "Complete"], "rows": qtbl_rows},
            "download": {"url": f"/api/ai/download/{fname}", "filename": fname},
            "suggestions": [f"Insights JAS 2026 vs JAS 2025 for {name}", f"Which product sold less in JAS 2026 vs JAS 2025 for {name}"]}
