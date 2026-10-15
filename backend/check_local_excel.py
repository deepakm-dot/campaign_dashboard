import openpyxl

path = r"c:\Users\deepak.m_myg\Downloads\All kerala SMS_2025 SEPT TO 2026 AUG STORE DATA.xlsx"
wb = openpyxl.load_workbook(path, read_only=True)
for name in wb.sheetnames:
    s = wb[name]
    count = 0
    sample = []
    for row in s.iter_rows(values_only=True):
        if row and row[0]:
            count += 1
            if len(sample) < 3:
                sample.append(row[0])
    print(f"Sheet: {name}, non-empty rows: {count}, sample: {sample}")
