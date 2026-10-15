import sqlite3
import openpyxl
import re
import os
import json
import base64
import urllib.request

DB_FILE = 'campaign_dashboard.db'
CH_URL  = "https://pdhsuv47ec.ap-south-1.aws.clickhouse.cloud:8443/"
CH_USER = "default"
CH_PASS = "ZFlujj9SA_Iei"

def query_clickhouse(sql):
    req = urllib.request.Request(CH_URL)
    req.add_header('Authorization', 'Basic ' + base64.b64encode(f"{CH_USER}:{CH_PASS}".encode()).decode())
    req.add_header('Content-Type', 'text/plain')
    req.data = (sql + " FORMAT JSON").encode()
    try:
        with urllib.request.urlopen(req, timeout=60) as r:
            return json.loads(r.read().decode()).get('data', [])
    except Exception as e:
        print(f"ClickHouse ERROR: {e}", flush=True)
        return []

def main():
    conn = sqlite3.connect(DB_FILE)
    conn.row_factory = sqlite3.Row
    c = conn.cursor()

    # 1. Check Row 67 and Row 37
    r67 = c.execute("SELECT * FROM campaigns WHERE sl_no=67").fetchone()
    r37 = c.execute("SELECT * FROM campaigns WHERE sl_no=37").fetchone()

    print(f"Current Row 67: {dict(r67) if r67 else None}")
    print(f"Current Row 37: {dict(r37) if r37 else None}")

    # Ensure Row 67 exists and has clean fields
    if not r67:
        c.execute("""
            INSERT INTO campaigns (campaign_name, type, start_date, end_date, source, campaign_group,
                                   data_source, budget, target_size, branch_code, sl_no, sheet_tab,
                                   purchase_count, purchase_value, conversion_rate, roi, new_customers, existing_customers)
            VALUES ('All kerala SMS', 'SMS', '2026-09-30', '2026-09-30', 'Vinu', 'Onam Climax September 30',
                    '2025 SEPT TO 2026 AUG STORE DATA', 168375.22, 1550592, '', 67, 'September 2026',
                    0, 0.0, 0.0, 0.0, 0, 0)
        """)
        camp_67_id = c.lastrowid
    else:
        camp_67_id = r67['id']
        c.execute("""
            UPDATE campaigns SET campaign_name='All kerala SMS', type='SMS', start_date='2026-09-30',
                                 end_date='2026-09-30', campaign_group='Onam Climax September 30',
                                 data_source='2025 SEPT TO 2026 AUG STORE DATA', target_size=1550592,
                                 sl_no=67, sheet_tab='September 2026'
            WHERE id=?
        """, (camp_67_id,))

    # Ensure Row 37 exists
    if not r37:
        c.execute("""
            INSERT INTO campaigns (campaign_name, type, start_date, end_date, source, campaign_group,
                                   data_source, budget, target_size, branch_code, sl_no, sheet_tab,
                                   purchase_count, purchase_value, conversion_rate, roi, new_customers, existing_customers)
            VALUES ('All kerala SMS', 'SMS', '2026-09-03', '2026-09-06', 'Vinu', 'Price Challenege September 3-6',
                    'Store Data Visited Last 3 year', 101986.95, 679913, '', 37, 'September 2026',
                    0, 0.0, 0.0, 0.0, 0, 0)
        """)
        camp_37_id = c.lastrowid
    else:
        camp_37_id = r37['id']

    conn.commit()
    print(f"Row 67 ID: {camp_67_id}, Row 37 ID: {camp_37_id}")

    # 2. Ingest customers for Row 67 from Downloads file
    c.execute("SELECT COUNT(*) FROM campaign_customers WHERE campaign_id=?", (camp_67_id,))
    count_67 = c.fetchone()[0]
    print(f"Campaign 67 current customer count: {count_67}")

    local_67_file = r"c:\Users\deepak.m_myg\Downloads\All kerala SMS_2025 SEPT TO 2026 AUG STORE DATA.xlsx"
    if count_67 == 0 and os.path.exists(local_67_file):
        print(f"Ingesting customers from {local_67_file} into campaign {camp_67_id}...")
        wb = openpyxl.load_workbook(local_67_file, read_only=True)
        batch = []
        total_inserted = 0
        for sname in wb.sheetnames:
            sheet = wb[sname]
            print(f"  Reading sheet: {sname}...")
            for row in sheet.iter_rows(values_only=True):
                if not row or not row[0]:
                    continue
                raw = str(row[0]).strip()
                clean = re.sub(r'\D', '', raw)
                if len(clean) >= 10:
                    mobile = clean[-10:]
                    batch.append((camp_67_id, mobile, '', ''))
                    if len(batch) >= 20000:
                        c.executemany("INSERT INTO campaign_customers (campaign_id, customer_mobile, customer_name, store_name) VALUES (?,?,?,?)", batch)
                        total_inserted += len(batch)
                        batch = []
        if batch:
            c.executemany("INSERT INTO campaign_customers (campaign_id, customer_mobile, customer_name, store_name) VALUES (?,?,?,?)", batch)
            total_inserted += len(batch)
        conn.commit()
        print(f"Successfully inserted {total_inserted} customers for campaign {camp_67_id}!")

    # 3. Check customer count for Row 37
    c.execute("SELECT COUNT(*) FROM campaign_customers WHERE campaign_id=?", (camp_37_id,))
    count_37 = c.fetchone()[0]
    print(f"Campaign 37 current customer count: {count_37}")
    if count_37 == 0:
        # Download Row 37 file from Drive ID 1Qtn0aQMjvXWNqDDmDEen_SWEXJ3Fe0zq
        dl_url = "https://drive.google.com/uc?export=download&id=1Qtn0aQMjvXWNqDDmDEen_SWEXJ3Fe0zq&confirm=t"
        dest_path = "all_kerala_sms_price_challenge.xlsx"
        if not os.path.exists(dest_path):
            print("Downloading Row 37 file from Google Drive...")
            req = urllib.request.Request(dl_url, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req, timeout=120) as r:
                content = r.read()
            if b'<html' in content[:200].lower():
                token_match = re.search(rb'confirm=([0-9A-Za-z_\-]+)', content)
                if token_match:
                    retry_url = f"https://drive.google.com/uc?export=download&id=1Qtn0aQMjvXWNqDDmDEen_SWEXJ3Fe0zq&confirm={token_match.group(1).decode()}"
                    with urllib.request.urlopen(urllib.request.Request(retry_url), timeout=120) as r2:
                        content = r2.read()
            with open(dest_path, 'wb') as f:
                f.write(content)
            print(f"Downloaded {len(content)/1024/1024:.2f} MB to {dest_path}")
        
        if os.path.exists(dest_path):
            print(f"Ingesting customers from {dest_path} into campaign {camp_37_id}...")
            wb = openpyxl.load_workbook(dest_path, read_only=True)
            batch = []
            total_inserted = 0
            for sname in wb.sheetnames:
                sheet = wb[sname]
                print(f"  Reading sheet: {sname}...")
                for row in sheet.iter_rows(values_only=True):
                    if not row or not row[0]:
                        continue
                    raw = str(row[0]).strip()
                    clean = re.sub(r'\D', '', raw)
                    if len(clean) >= 10:
                        mobile = clean[-10:]
                        batch.append((camp_37_id, mobile, '', ''))
                        if len(batch) >= 20000:
                            c.executemany("INSERT INTO campaign_customers (campaign_id, customer_mobile, customer_name, store_name) VALUES (?,?,?,?)", batch)
                            total_inserted += len(batch)
                            batch = []
            if batch:
                c.executemany("INSERT INTO campaign_customers (campaign_id, customer_mobile, customer_name, store_name) VALUES (?,?,?,?)", batch)
                total_inserted += len(batch)
            conn.commit()
            print(f"Successfully inserted {total_inserted} customers for campaign {camp_37_id}!")

    conn.close()

if __name__ == '__main__':
    main()
